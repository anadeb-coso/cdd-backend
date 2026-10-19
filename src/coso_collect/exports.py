"""Exports des données d'un formulaire COSOCollect (même esprit que les téléchargements kobocollect).

- XLSX : une feuille principale (une ligne par saisie, colonnes `_index`, métadonnées puis champs) et
  une feuille par groupe répétable (`_index`, `_parent_index`, `_parent_table_name`, `_submission__uuid`).
- CSV : feuille principale seulement.
- GeoJSON : un point par coordonnée GPS saisie.
- ZIP : toutes les photos / documents joints, rangés par saisie.
"""
import csv
import io
import json
import re
import tempfile
import zipfile

from django.utils import timezone

from .fields import CONTAINER_TYPES, SELECT_TYPES, form_fields
from .models import CollectSubmission
from .titles import parse_iso

META_COLUMNS = [
    ('_index', '_index'),
    ('_id', '_id'),
    ('_uuid', '_uuid'),
    ('_title', 'titre_combine'),
    ('_status', '_status'),
    ('_author', '_submitted_by'),
    ('_started', '_start'),
    ('_submitted', '_submission_time'),
    ('_source', '_source'),
    ('_version', '_version'),
]

DEFAULT_OPTIONS = {
    'labels': 'labels',        # valeurs des choix : 'labels' (libellés) ou 'values' (valeurs XML)
    'headers': 'labels',       # en-têtes : 'labels' ou 'names'
    'group_headers': True,     # préfixer l'en-tête par le chemin de groupes
    'group_sep': '/',
    'split_multiple': False,   # select_multiple : une colonne par choix (1/0)
    'media_urls': True,        # colonne « _URL » pour les photos / documents
}


def normalize_options(raw):
    raw = raw or {}
    out = dict(DEFAULT_OPTIONS)
    if raw.get('labels') in ('labels', 'values'):
        out['labels'] = raw['labels']
    if raw.get('headers') in ('labels', 'names'):
        out['headers'] = raw['headers']
    for key in ('group_headers', 'split_multiple', 'media_urls'):
        if key in raw:
            out[key] = raw[key] in (True, 'true', '1', 'on', 'yes')
    sep = raw.get('group_sep')
    if sep in ('/', '.', '-', ':', '_'):
        out['group_sep'] = sep
    return out


def _key_path(field):
    return field['local_path'].replace('.', '/')


def header_for(field, options, suffix=''):
    if options['headers'] == 'names':
        base = _key_path(field).replace('/', options['group_sep']) if options['group_headers'] else field['key']
    else:
        labels = field['group_labels'] + [field['label']] if options['group_headers'] else [field['label']]
        base = options['group_sep'].join(labels)
    return base + suffix


def _fmt_date(value, kind):
    dt = parse_iso(value)
    if not dt:
        return value
    return dt.strftime({'date': '%Y-%m-%d', 'datetime': '%Y-%m-%d %H:%M:%S', 'time': '%H:%M'}[kind])


class Table:
    def __init__(self, name, title):
        self.name = name
        self.title = title
        self.columns = []   # [(col_id, header)]
        self._known = set()
        self.rows = []

    def add_column(self, col_id, header):
        if col_id not in self._known:
            self._known.add(col_id)
            self.columns.append((col_id, header))


def _sheet_title(name, used):
    base = re.sub(r'[\[\]\*\?/\\:]', '_', name)[:28] or 'repeat'
    title, i = base, 2
    while title.lower() in used:
        title = f'{base[:25]}_{i}'
        i += 1
    used.add(title.lower())
    return title


def build_tables(form, submissions, options, attachment_url=None):
    """Tables (principale + répétables) prêtes à écrire, colonnes dans l'ordre du formulaire."""
    options = normalize_options(options)
    fields = form_fields(form.form)
    used_titles = {'data'}
    main = Table('main', 'data')
    for col_id, header in META_COLUMNS:
        main.add_column(col_id, header)
    repeat_tables = {}

    def declare(field_list, table):
        for f in field_list:
            if f['type'] == 'note':
                continue
            if f['type'] == 'group':
                declare(f['children'], table)
            elif f['type'] == 'repeat':
                sub = Table(f['path'], _sheet_title(f['key'], used_titles))
                for col_id in ('_index', '_parent_index', '_parent_table_name', '_submission__id', '_submission__uuid'):
                    sub.add_column(col_id, col_id)
                repeat_tables[f['path']] = sub
                declare(f['children'], sub)
            else:
                table.add_column(f['path'], header_for(f, options))
                if f['type'] in ('select_multiple', 'select_multiple_check') and options['split_multiple']:
                    for value, label in f['choices'].items():
                        choice_header = label if options['labels'] == 'labels' else value
                        table.add_column(f'{f["path"]}::{value}', header_for(f, options, '/' + choice_header))
                if f['type'] == 'geopoint':
                    for part in ('latitude', 'longitude', 'altitude', 'precision'):
                        table.add_column(f'{f["path"]}::{part}', '_' + header_for(f, options) + '_' + part)
                if f['type'] in ('photo', 'file') and options['media_urls']:
                    table.add_column(f'{f["path"]}::url', header_for(f, options) + '_URL')

    declare(fields, main)

    def put(row, f, value):
        if value is None or value == '' or value == []:
            return
        ftype = f['type']
        if ftype == 'select_one':
            row[f['path']] = f['choices'].get(str(value), str(value)) if options['labels'] == 'labels' else value
        elif ftype in ('select_multiple', 'select_multiple_check'):
            values = value if isinstance(value, list) else [value]
            if options['labels'] == 'labels':
                row[f['path']] = ', '.join(f['choices'].get(str(v), str(v)) for v in values)
            else:
                row[f['path']] = ' '.join(str(v) for v in values)
            if options['split_multiple']:
                chosen = {str(v) for v in values}
                for choice in f['choices']:
                    row[f'{f["path"]}::{choice}'] = 1 if choice in chosen else 0
        elif ftype in ('date', 'datetime', 'time'):
            row[f['path']] = _fmt_date(value, ftype)
        elif ftype == 'geopoint' and isinstance(value, dict):
            parts = [value.get('latitude'), value.get('longitude'), value.get('altitude'), value.get('accuracy')]
            row[f['path']] = ' '.join('' if p is None else str(p) for p in parts).strip()
            for name, part in zip(('latitude', 'longitude', 'altitude', 'precision'), parts):
                row[f'{f["path"]}::{name}'] = part
        elif ftype in ('photo', 'file') and isinstance(value, dict):
            row[f['path']] = value.get('name') or ''
            if options['media_urls'] and attachment_url and value.get('attachment_id'):
                row[f'{f["path"]}::url'] = attachment_url(value['attachment_id'])
        elif isinstance(value, (dict, list)):
            row[f['path']] = json.dumps(value, ensure_ascii=False)
        else:
            row[f['path']] = value

    def fill(field_list, values, row, submission, parent_table, parent_index):
        for f in field_list:
            if f['type'] == 'note' or not isinstance(values, dict):
                continue
            value = values.get(f['key'])
            if f['type'] == 'group':
                fill(f['children'], value or {}, row, submission, parent_table, parent_index)
            elif f['type'] == 'repeat':
                table = repeat_tables[f['path']]
                for item in value if isinstance(value, list) else []:
                    sub_row = {
                        '_index': len(table.rows) + 1, '_parent_index': parent_index,
                        '_parent_table_name': parent_table.title, '_submission__id': submission.pk,
                        '_submission__uuid': str(submission.uuid),
                    }
                    table.rows.append(sub_row)
                    fill(f['children'], item or {}, sub_row, submission, table, sub_row['_index'])
            else:
                put(row, f, value)

    status_labels = dict(CollectSubmission.STATUSES)
    by_page = {}
    for f in fields:
        by_page.setdefault(f['page'], []).append(f)
    for submission in submissions:
        row = {
            '_index': len(main.rows) + 1,
            '_id': submission.pk,
            '_uuid': str(submission.uuid),
            '_title': submission.title,
            '_status': str(status_labels.get(submission.status, submission.status)),
            '_author': submission.author_name,
            '_started': _local(submission.started_at),
            '_submitted': _local(submission.submitted_at),
            '_source': submission.source,
            '_version': submission.form_version,
        }
        main.rows.append(row)
        data = submission.data if isinstance(submission.data, list) else []
        for page_index, page_fields in by_page.items():
            values = data[page_index] if page_index < len(data) else {}
            fill(page_fields, values or {}, row, submission, main, row['_index'])
    return [main] + list(repeat_tables.values())


def _local(dt):
    if not dt:
        return None
    return timezone.localtime(dt).strftime('%Y-%m-%d %H:%M:%S')


def to_xlsx(form, submissions, options, attachment_url=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    tables = build_tables(form, submissions, options, attachment_url)
    wb = Workbook()
    wb.remove(wb.active)
    for table in tables:
        ws = wb.create_sheet(table.title)
        ws.append([header for _col, header in table.columns])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in table.rows:
            ws.append([_cell(row.get(col_id)) for col_id, _h in table.columns])
        ws.freeze_panes = 'A2'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _cell(value):
    if isinstance(value, str) and len(value) > 32000:
        return value[:32000]
    return value


def to_csv(form, submissions, options, attachment_url=None):
    main = build_tables(form, submissions, options, attachment_url)[0]
    out = io.StringIO()
    writer = csv.writer(out, delimiter=';')
    writer.writerow([header for _col, header in main.columns])
    for row in main.rows:
        writer.writerow(['' if row.get(col_id) is None else row.get(col_id) for col_id, _h in main.columns])
    return ('﻿' + out.getvalue()).encode('utf-8')


def to_geojson(form, submissions, point_rows):
    """`point_rows` : itérable de (submission, path, field, value) — cf. fields.geopoint_values."""
    features = []
    for submission, path, field, value in point_rows:
        features.append({
            'type': 'Feature',
            'geometry': {'type': 'Point', 'coordinates': [value['longitude'], value['latitude']]},
            'properties': {
                'title': submission.title, 'uuid': str(submission.uuid), 'status': submission.status,
                'author': submission.author_name, 'field': field['label'], 'path': path,
                'accuracy': value.get('accuracy'), 'altitude': value.get('altitude'),
            },
        })
    return json.dumps({'type': 'FeatureCollection', 'name': form.title, 'features': features},
                      ensure_ascii=False).encode('utf-8')


def _safe_name(text, default='fichier'):
    text = re.sub(r'[\\/:*?"<>|\r\n]+', '_', str(text or '')).strip(' .')
    return text[:120] or default


def zip_attachments(attachments):
    """Archive ZIP (fichier temporaire, prêt pour FileResponse) des pièces jointes, rangées par saisie."""
    tmp = tempfile.TemporaryFile()
    used = set()
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as archive:
        for attachment in attachments:
            folder = _safe_name(attachment.submission.title if attachment.submission else attachment.submission_uuid,
                                str(attachment.submission_uuid))
            field = _safe_name(attachment.field_path.split('.')[-1] if attachment.field_path else '', 'piece')
            name = f'{folder}/{field}_{_safe_name(attachment.name)}'
            base, i = name, 2
            while name in used:
                name = f'{base.rsplit(".", 1)[0]}_{i}' + (f'.{base.rsplit(".", 1)[1]}' if '.' in base else '')
                i += 1
            used.add(name)
            try:
                with attachment.file.open('rb') as handle:
                    archive.writestr(name, handle.read())
            except Exception:  # fichier absent du stockage : on continue avec les autres
                continue
    tmp.seek(0)
    return tmp


def select_fields_for_reports(form):
    """Champs « rapport » : choix (fréquences) et nombres (statistiques)."""
    out = []

    def walk(field_list):
        for f in field_list:
            if f['type'] in CONTAINER_TYPES:
                walk(f['children'])
            elif f['type'] in SELECT_TYPES or f['type'] in ('integer', 'decimal'):
                out.append(f)

    walk(form_fields(form.form))
    return out
