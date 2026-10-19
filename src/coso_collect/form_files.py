"""Import / export du fichier d'un formulaire COSOCollect.

- `.xlsx` : XLSForm (même conversion que le Générateur de formulaire des tâches, cf.
  dashboard/process_manager/tasks/form_design.py), le titre combiné étant porté par la colonne standard
  `instance_name` de la feuille `settings` (équivalent kobocollect), plus des colonnes propres
  (`cosocollect_separator`, `cosocollect_slugify`, `cosocollect_single_submission`) ignorées par les autres
  outils.
- `.json` : format natif, sans perte ({"format": "cosocollect", "form": [...], "title_template": {...},
  "single_submission": false}).
"""
import io
import json
import re

from dashboard.process_manager.tasks.form_design import (
    form_design_to_xlsform, validate_form_design, xlsform_to_form_design,
)

from .fields import leaf_fields
from .titles import DEFAULT_SEPARATOR, normalize_template

JSON_FORMAT = 'cosocollect'
JSON_VERSION = 1

_META_EXPR = {
    'created_at': "format-date-time(now(), '%Y%m%d%H%M%S')",
    'created_date': "format-date(today(), '%Y%m%d')",
    'created_time': "format-date-time(now(), '%H%M%S')",
    'author': '${username}',
}


def _quote(text):
    return "'" + str(text).replace("'", "") + "'"


def template_to_instance_name(form, title_template):
    template = normalize_template(title_template, form.form)
    by_path = {f['path']: f for f in leaf_fields(form.form)}
    args = []
    for part in template['parts']:
        if args and template['separator']:
            args.append(_quote(template['separator']))
        if part['type'] == 'field':
            field = by_path.get(part['path'])
            if field:
                args.append('${' + field['key'] + '}')
        elif part['type'] == 'meta':
            args.append(_META_EXPR.get(part['key']) or _quote(form.title))
        else:
            args.append(_quote(part['value']))
    if not args:
        return ''
    return args[0] if len(args) == 1 else 'concat(' + ', '.join(args) + ')'


def _split_args(expr):
    args, depth, quote, cur = [], 0, None, ''
    for ch in expr:
        if quote:
            cur += ch
            if ch == quote:
                quote = None
            continue
        if ch in ("'", '"'):
            quote = ch
            cur += ch
        elif ch == '(':
            depth += 1
            cur += ch
        elif ch == ')':
            depth -= 1
            cur += ch
        elif ch == ',' and depth == 0:
            args.append(cur.strip())
            cur = ''
        else:
            cur += ch
    if cur.strip():
        args.append(cur.strip())
    return args


def instance_name_to_template(instance_name, form_pages, separator=None, slugify=True):
    """Expression XLSForm `instance_name` -> `title_template` (meilleur effort)."""
    expr = (instance_name or '').strip()
    if not expr:
        return {'parts': [], 'separator': separator if separator is not None else DEFAULT_SEPARATOR,
                'slugify': slugify}
    match = re.match(r'^concat\s*\((.*)\)\s*$', expr, re.S)
    args = _split_args(match.group(1)) if match else [expr]
    literals = [a[1:-1] for a in args if len(a) >= 2 and a[0] == a[-1] and a[0] in "'\""]
    if separator is None:
        separator = max(set(literals), key=literals.count) if literals else DEFAULT_SEPARATOR
    by_key = {}
    for field in leaf_fields(form_pages):
        if not field['repeat_path']:
            by_key.setdefault(field['key'], field['path'])

    parts = []
    for arg in args:
        ref = re.fullmatch(r'\$\{([^}]+)\}', arg)
        if ref:
            name = ref.group(1).strip()
            if name == 'username':
                parts.append({'type': 'meta', 'key': 'author'})
            elif name in by_key:
                parts.append({'type': 'field', 'path': by_key[name]})
        elif 'now()' in arg or 'today()' in arg:
            if '%H' in arg and '%Y' in arg:
                key = 'created_at'
            elif '%Y' in arg or 'today()' in arg:
                key = 'created_date'
            else:
                key = 'created_time'
            parts.append({'type': 'meta', 'key': key})
        elif len(arg) >= 2 and arg[0] == arg[-1] and arg[0] in "'\"":
            text = arg[1:-1]
            if text and text != separator:
                parts.append({'type': 'text', 'value': text})
    return normalize_template({'parts': parts, 'separator': separator, 'slugify': slugify}, form_pages)


def export_xlsform(form):
    from openpyxl import load_workbook

    template = normalize_template(form.title_template, form.form)
    data = form_design_to_xlsform(form.form or [], form_title=form.title,
                                  instance_name=template_to_instance_name(form, template) or None)
    extra = []
    if template['parts']:
        extra += [('cosocollect_separator', template['separator']),
                  ('cosocollect_slugify', 'yes' if template['slugify'] else 'no')]
    if form.single_submission:
        extra.append(('cosocollect_single_submission', 'yes'))
    if not extra:
        return data
    wb = load_workbook(io.BytesIO(data))
    sheet = wb['settings']
    col = sheet.max_column
    for offset, (name, value) in enumerate(extra, start=1):
        sheet.cell(row=1, column=col + offset, value=name)
        sheet.cell(row=2, column=col + offset, value=value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_json(form):
    payload = {
        'format': JSON_FORMAT,
        'version': JSON_VERSION,
        'title': form.title,
        'description': form.description,
        'form': form.form or [],
        'title_template': normalize_template(form.title_template, form.form),
        'single_submission': form.single_submission,
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode('utf-8')


def _yes(value):
    return str(value or '').strip().lower() in ('yes', 'true', '1', 'oui')


def _read_settings(raw_bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    if 'settings' not in wb.sheetnames:
        return {}
    rows = list(wb['settings'].iter_rows(values_only=True))
    if len(rows) < 2:
        return {}
    head = [str(c).strip() if c is not None else '' for c in rows[0]]
    return {h: rows[1][i] for i, h in enumerate(head) if h and i < len(rows[1]) and rows[1][i] is not None}


def import_form_file(uploaded):
    """Lit un fichier .json ou .xlsx -> dict {title, description, form, title_template, single_submission,
    errors}."""
    name = (getattr(uploaded, 'name', '') or '').lower()
    raw = uploaded.read()
    if name.endswith('.json'):
        try:
            payload = json.loads(raw.decode('utf-8'))
        except (ValueError, UnicodeDecodeError):
            return {'errors': ['Fichier JSON illisible.']}
        if isinstance(payload, list):  # tableau de pages brut (format Task.form)
            payload = {'form': payload}
        if not isinstance(payload, dict):
            return {'errors': ['Format JSON inattendu.']}
        pages = payload.get('form') or []
        clean, errors = validate_form_design(pages)
        pages = clean if clean is not None else []
        return {
            'title': str(payload.get('title') or '').strip()[:255],
            'description': str(payload.get('description') or ''),
            'form': pages,
            'title_template': normalize_template(payload.get('title_template'), pages),
            'single_submission': payload.get('single_submission') is True,
            'errors': errors,
        }
    if name.endswith('.xlsx'):
        try:
            pages = xlsform_to_form_design(io.BytesIO(raw))
            settings = _read_settings(raw)
        except Exception as exc:  # classeur illisible / feuille `survey` absente
            return {'errors': ['Fichier XLSForm illisible : {}'.format(exc)]}
        clean, errors = validate_form_design(pages)
        pages = clean if clean is not None else pages
        separator = settings.get('cosocollect_separator')
        slugify = str(settings.get('cosocollect_slugify', 'yes')).strip().lower() not in ('no', 'false', '0')
        return {
            'title': str(settings.get('form_title') or '').strip()[:255],
            'description': '',
            'form': pages,
            'title_template': instance_name_to_template(
                settings.get('instance_name'), pages,
                separator=str(separator) if separator is not None else None, slugify=slugify,
            ),
            'single_submission': _yes(settings.get('cosocollect_single_submission')),
            'errors': errors,
        }
    return {'errors': ['Format non pris en charge (.xlsx ou .json attendu).']}
