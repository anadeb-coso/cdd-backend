"""Imports Excel du DCC (aperçu puis confirmation, comme la page « Localités ») :

  - comptes de facilitateurs (canevas `ac_list_canevas.xlsx` : Nom, Prénom, Sex, Numéro de téléphone,
    adresse mail, Facilitateur FC/FT), créés comme par le formulaire (dashboard/facilitators/creation.py) ;
  - localités initiales (villages d'affectation) ou de stabilisation des facilitateurs (canevas
    `ac_affectation_villages.xlsx` : une ligne par village, avec le FC et le FT qui le suivent), enregistrées
    comme par la page « Localités » (dashboard/facilitators/localities_apply.py).

Les lignes lues sont transmises de l'aperçu à la confirmation dans un champ signé (rien n'est gardé côté
serveur) ; tout est recontrôlé à la confirmation."""
import difflib
import re
import unicodedata
from collections import OrderedDict

from django.core import signing
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.utils.translation import gettext as _

from authentication.accounts import email_in_use, generate_username, username_candidates

SIGNING_SALT = 'dashboard.facilitators.imports'
SIGNING_MAX_AGE = 3 * 3600
MAX_FACILITATOR_ROWS = 150
MAX_ASSIGNMENT_ROWS = 5000

FACILITATOR_TYPES_BY_CODE = {
    'FC': 'community_facilitator', 'COMMUNITYFACILITATOR': 'community_facilitator',
    'FACILITATEURCOMMUNAUTAIRE': 'community_facilitator',
    'FT': 'technical_facilitator', 'TECHNICALFACILITATOR': 'technical_facilitator',
    'FACILITATEURTECHNIQUE': 'technical_facilitator',
}
SEX_BY_CODE = {'M': 'M.', 'H': 'M.', 'HOMME': 'M.', 'MASCULIN': 'M.', 'MR': 'M.',
               'F': 'Mme', 'FEMME': 'Mme', 'FEMININ': 'Mme', 'MME': 'Mme', 'MLLE': 'Mme'}

FACILITATOR_COLUMNS = OrderedDict([
    ('last_name', ('Nom', ('nom', 'noms', 'last name'))),
    ('first_name', ('Prénom', ('prenom', 'prenoms', 'first name', 'first names'))),
    ('sex', ('Sex', ('sex', 'sexe'))),
    ('phone', ('Numéro de téléphone', ('numero de telephone', 'telephone', 'contact', 'phone', 'numero'))),
    ('email', ('adresse mail', ('adresse mail', 'mail', 'email', 'e-mail', 'adresse e-mail', 'adresse email'))),
    ('facilitator_type', ('Facilitateur', ('facilitateur', 'type', 'type de facilitateur'))),
])

ASSIGNMENT_COLUMNS = OrderedDict([
    ('region', ('Région', ('region',))),
    ('prefecture', ('Préfecture', ('prefecture',))),
    ('commune', ('Commune', ('commune',))),
    ('canton', ('Canton', ('canton',))),
    ('village', ('Village/Quartier', ('village/quartier', 'village', 'quartier', 'village / quartier'))),
    ('cvd', ('CVD', ('cvd',))),
    ('headquarters', ('Siège CVD', ('siege cvd', 'siege du cvd'))),
    ('geographical_unit', ('Unité géographique', ('unite geographique', 'ug'))),
    ('fc_name', ('Nom de FC', ('nom de fc', 'nom du fc', 'fc'))),
    ('fc_phone', ('Contact de FC', ('contact de fc', 'contact du fc', 'telephone de fc'))),
    ('fc_email', ('Adresse mail de FC', ('adresse mail de fc', 'email de fc', 'mail de fc', 'adresse mail du fc'))),
    ('ft_name', ('Nom de FT', ('nom de ft', 'nom du ft', 'ft'))),
    ('ft_phone', ('Contact de FT', ('contact de ft', 'contact du ft', 'telephone de ft'))),
    ('ft_email', ('Adresse mail de FT', ('adresse mail de ft', 'email de ft', 'mail de ft', 'adresse mail du ft'))),
])

MODE_INITIAL = 'initial'
MODE_STABILIZATION = 'stabilization'


class ImportFileError(Exception):
    pass


# ---------------------------------------------------------------------------------------------
# Lecture du fichier
# ---------------------------------------------------------------------------------------------

def normalized(value):
    """Clé de comparaison : majuscules sans accents, lettres et chiffres seulement (comme
    dashboard/statistics/utils.py::normaliser_chaine)."""
    value = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii').upper()
    return re.sub(r'[^A-Z0-9]', '', value)


def _header_key(value):
    value = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'\s+', ' ', value).strip()


def cell_text(value):
    """Texte d'une cellule : nombres entiers sans « .0 » (téléphones), espaces superflus retirés."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r'\s+', ' ', str(value)).strip()


def read_sheet(uploaded_file, columns, max_rows):
    """Lignes (numéro de ligne Excel, {clé: texte}) de la 1re feuille ; colonnes reconnues par leur intitulé
    (sans tenir compte de la casse ni des accents). Lève ImportFileError si le fichier est illisible ou s'il
    manque une colonne."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception:
        raise ImportFileError(_("The file could not be read: please use an Excel file (.xlsx) built from the template."))
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    positions = {}
    header_keys = [_header_key(h) for h in header]
    for key, (label, aliases) in columns.items():
        for index, header_key in enumerate(header_keys):
            if header_key in aliases and index not in positions.values():
                positions[key] = index
                break
    missing = [label for key, (label, _aliases) in columns.items() if key not in positions]
    if missing:
        workbook.close()
        raise ImportFileError(_("Missing column(s) in the file: %(columns)s.") % {'columns': ", ".join(missing)})

    lines = []
    for number, values in enumerate(rows, start=2):
        data = {key: cell_text(values[index]) if index < len(values) else '' for key, index in positions.items()}
        if not any(data.values()):
            continue
        lines.append((number, data))
        if len(lines) > max_rows:
            workbook.close()
            raise ImportFileError(_("Too many rows (more than %(n)s): please split the file.") % {'n': max_rows})
    workbook.close()
    if not lines:
        raise ImportFileError(_("The file contains no data row."))
    return lines


def template_workbook(columns, examples=()):
    """Canevas vide (intitulés du canevas fourni), pour le bouton « Télécharger le modèle »."""
    import openpyxl
    from openpyxl.styles import Font

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Feuil1'
    sheet.append([label for label, _aliases in columns.values()])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for example in examples:
        sheet.append(list(example))
    for index, (label, _aliases) in enumerate(columns.values(), start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = max(14, len(label) + 4)
    return workbook


def sign(payload):
    return signing.dumps(payload, salt=SIGNING_SALT, compress=True)


def unsign(token):
    """Contenu signé de l'aperçu, ou None (modifié, expiré ou absent)."""
    try:
        return signing.loads(token or '', salt=SIGNING_SALT, max_age=SIGNING_MAX_AGE)
    except signing.BadSignature:
        return None


# ---------------------------------------------------------------------------------------------
# Import des comptes de facilitateurs
# ---------------------------------------------------------------------------------------------

def facilitator_type_of(value):
    return FACILITATOR_TYPES_BY_CODE.get(normalized(value))


def sex_of(value):
    return SEX_BY_CODE.get(normalized(value)) if value else None


def check_facilitator_rows(lines):
    """Contrôle des lignes lues : [{line, last_name, first_name, sex, phone, email, facilitator_type, username,
    errors, warnings}] ; une ligne en erreur n'est pas créée. Les identifiants proposés tiennent compte des
    lignes précédentes du fichier."""
    from authentication.models import Facilitator

    existing_names = {normalized(name) for name in Facilitator.objects.values_list('name', flat=True) if name}
    seen_emails, seen_names, reserved = set(), set(), set()
    rows = []
    for number, data in lines:
        row = {
            'line': number,
            'last_name': data['last_name'].upper(),
            'first_name': data['first_name'],
            'sex': sex_of(data['sex']),
            'phone': data['phone'] or None,
            'email': data['email'] or None,
            'facilitator_type': facilitator_type_of(data['facilitator_type']),
            'facilitator_type_code': data['facilitator_type'],
            'username': None,
            'errors': [],
            'warnings': [],
        }
        if not row['last_name'] or not row['first_name'] or not username_candidates(row['first_name'], row['last_name']):
            row['errors'].append(_("The last name and the first name are required."))
        if not row['facilitator_type']:
            row['errors'].append(_("Unknown facilitator type « %(value)s » (expected FC or FT).") % {'value': data['facilitator_type']})
        if data['sex'] and not row['sex']:
            row['warnings'].append(_("Unknown sex « %(value)s » (expected M or F): left empty.") % {'value': data['sex']})
        if row['email']:
            try:
                validate_email(row['email'])
            except ValidationError:
                row['errors'].append(_("Invalid e-mail address."))
            else:
                if row['email'].lower() in seen_emails:
                    row['errors'].append(_("This e-mail address is already used by a previous row of the file."))
                elif email_in_use(row['email']):
                    row['errors'].append(_("An account (user or facilitator) already uses this e-mail address."))
        else:
            row['warnings'].append(_("No e-mail address: no MGP account, the login details will be displayed after the creation."))
        name_key = normalized(f"{row['last_name']} {row['first_name']}")
        if name_key and name_key in existing_names:
            row['errors'].append(_("A facilitator with the same name already exists."))
        elif name_key and name_key in seen_names:
            row['errors'].append(_("Same name as a previous row of the file."))

        if not row['errors']:
            row['username'] = generate_username(row['first_name'], row['last_name'], reserved=reserved)
            reserved.add(row['username'])
            seen_names.add(name_key)
            if row['email']:
                seen_emails.add(row['email'].lower())
        rows.append(row)
    return rows


# ---------------------------------------------------------------------------------------------
# Import des localités (initiales ou de stabilisation)
# ---------------------------------------------------------------------------------------------

class LocalityIndex:
    """Recherche des villages par leur nom et ceux de leurs parents (région, préfecture, commune, canton), avec
    repli sur le CVD ; noms comparés sans accents, casse, espaces ni ponctuation."""

    LEVELS = ('region', 'prefecture', 'commune', 'canton')
    TYPES = {'region': 'Region', 'prefecture': 'Prefecture', 'commune': 'Commune', 'canton': 'Canton'}

    def __init__(self, tree):
        from administrativelevels.models import CVD
        from cdd.call_objects_from_other_db import mis_objects_call

        self.tree = tree
        self.villages_by_name = {}
        self.names_by_parent = {}
        for _id, name, _type, parent_id, _cvd_id in tree.levels.values():
            if _type == 'Village':
                self.villages_by_name.setdefault(normalized(name), []).append(_id)
                self.names_by_parent.setdefault(parent_id, []).append(name)
        self.cvds_by_name = {}
        self.cvd_headquarters = {}
        for cvd_id, name, headquarters_id in mis_objects_call.get_all_objects(CVD).values_list('id', 'name', 'headquarters_village_id'):
            self.cvds_by_name.setdefault(normalized(name), []).append(cvd_id)
            self.cvd_headquarters[cvd_id] = headquarters_id

    def ancestors(self, village_id):
        """{type: nom normalisé} des parents d'un village."""
        result = {}
        level = self.tree.levels.get(village_id)
        parent_id = level[3] if level else None
        while parent_id:
            parent = self.tree.levels.get(parent_id)
            if not parent:
                break
            result[parent[2]] = normalized(parent[1])
            parent_id = parent[3]
        return result

    def _best(self, candidates, data):
        """Candidats dont le plus de parents portent le nom donné (au moins un si des parents sont donnés)."""
        given = {self.TYPES[level]: normalized(data.get(level)) for level in self.LEVELS if normalized(data.get(level))}
        scored = []
        for village_id in candidates:
            ancestors = self.ancestors(village_id)
            scored.append((sum(1 for _type, name in given.items() if ancestors.get(_type) == name), village_id))
        if not scored:
            return [], 0, len(given)
        best = max(score for score, _id in scored)
        if given and best == 0:
            return [], 0, len(given)
        return [village_id for score, village_id in scored if score == best], best, len(given)

    def resolve(self, data):
        """(id du village | None, avertissements, erreurs)."""
        warnings, errors = [], []
        village = data.get('village')
        candidates, score, given = self._best(self.villages_by_name.get(normalized(village), []), data)
        if len(candidates) == 1:
            if score < given:
                warnings.append(_("Village « %(name)s » found, but some of its parent levels have another name in the DCC.") % {'name': village})
            return candidates[0], warnings, errors
        if len(candidates) > 1:
            errors.append(_("Several villages « %(name)s » match this row: please check the region, prefecture, commune and canton.") % {'name': village})
            return None, warnings, errors

        # Village introuvable : le CVD (ses villages sont de toute façon tous affectés ensemble).
        cvd_villages = []
        for cvd_id in self.cvds_by_name.get(normalized(data.get('cvd')), []):
            members = [v for v in self.tree.cvd_members.get(cvd_id, [])]
            headquarters = self.cvd_headquarters.get(cvd_id)
            if headquarters in members:
                members = [headquarters] + [m for m in members if m != headquarters]
            cvd_villages.append(members)
        matching = []
        for members in cvd_villages:
            found, _score, _given = self._best(members, data)
            if found:
                matching.append(members[0])
        if len(matching) == 1:
            warnings.append(_("Village « %(village)s » not found in the DCC: the villages of the CVD « %(cvd)s » are used.") % {
                'village': village, 'cvd': data.get('cvd')})
            return matching[0], warnings, errors

        suggestions = self.suggestions(data)
        message = _("Village « %(name)s » not found in the DCC.") % {'name': village}
        if suggestions:
            message += " " + _("Did you mean: %(names)s?") % {'names': ", ".join(suggestions)}
        errors.append(message)
        return None, warnings, errors

    def suggestions(self, data):
        """Noms de villages proches, sous le canton donné s'il est trouvé, sinon dans tout le pays."""
        names = []
        canton_key = normalized(data.get('canton'))
        if canton_key:
            for _id, name, _type, _parent, _cvd in self.tree.levels.values():
                if _type == 'Canton' and normalized(name) == canton_key:
                    names += self.names_by_parent.get(_id, [])
        if not names:
            names = [self.tree.levels[ids[0]][1] for ids in self.villages_by_name.values()]
        by_key = {normalized(name): name for name in names}
        return [by_key[key] for key in difflib.get_close_matches(normalized(data.get('village')), list(by_key), n=3, cutoff=0.6)]


def phone_key(value):
    digits = re.sub(r'\D', '', str(value or ''))
    return digits[-8:] if len(digits) >= 8 else None


class FacilitatorIndex:
    """Facilitateurs du DCC recherchés par e-mail, puis téléphone et nom."""

    def __init__(self):
        from authentication.models import Facilitator

        self.facilitators = list(Facilitator.objects.only(
            'id', 'name', 'email', 'phone', 'facilitator_type', 'active', 'username'))
        self.by_email, self.by_phone, self.by_name = {}, {}, {}
        for facilitator in self.facilitators:
            if facilitator.email:
                self.by_email.setdefault(facilitator.email.strip().lower(), []).append(facilitator)
            if phone_key(facilitator.phone):
                self.by_phone.setdefault(phone_key(facilitator.phone), []).append(facilitator)
            if facilitator.name:
                self.by_name.setdefault(normalized(facilitator.name), []).append(facilitator)

    def resolve(self, name, phone, email):
        """(facilitateur | None, avertissements, erreurs)."""
        warnings, errors = [], []
        label = name or email or phone
        found = self.by_email.get((email or '').strip().lower(), []) if email else []
        if not found:
            by_phone = self.by_phone.get(phone_key(phone), []) if phone_key(phone) else []
            by_name = self.by_name.get(normalized(name), []) if normalized(name) else []
            both = [f for f in by_phone if f in by_name]
            found = both or (by_name if len(by_name) == 1 else []) or (by_phone if len(by_phone) == 1 and not name else [])
            if found and email:
                warnings.append(_("Facilitator « %(label)s » found by name/phone, not by the e-mail address %(email)s.") % {
                    'label': label, 'email': email})
        if not found:
            errors.append(_("Facilitator « %(label)s » not found in the DCC (create the account first).") % {'label': label})
            return None, warnings, errors
        if len(found) > 1:
            errors.append(_("Several facilitators match « %(label)s ».") % {'label': label})
            return None, warnings, errors
        facilitator = found[0]
        if not facilitator.active:
            errors.append(_("The facilitator « %(label)s » is inactive.") % {'label': label})
            return None, warnings, errors
        return facilitator, warnings, errors


def group_assignment_rows(lines, locality_index, facilitator_index):
    """Lignes du fichier -> ({id facilitateur: {'facilitator', 'roles', 'villages', 'lines'}}, lignes en erreur,
    avertissements par ligne). Un village est donné au FC et au FT de sa ligne."""
    by_facilitator = OrderedDict()
    line_errors, line_warnings = [], []
    for number, data in lines:
        village_id, warnings, errors = locality_index.resolve(data)
        line_warnings += [(number, w) for w in warnings]
        if errors:
            line_errors += [(number, e) for e in errors]
            continue
        people = 0
        for role, expected_type in (('fc', 'community_facilitator'), ('ft', 'technical_facilitator')):
            name, phone, email = data[f'{role}_name'], data[f'{role}_phone'], data[f'{role}_email']
            if not (name or phone or email):
                continue
            people += 1
            facilitator, warnings, errors = facilitator_index.resolve(name, phone, email)
            line_warnings += [(number, w) for w in warnings]
            if errors:
                line_errors += [(number, e) for e in errors]
                continue
            if facilitator.facilitator_type != expected_type:
                line_warnings.append((number, _("« %(name)s » is not registered as a %(role)s in the DCC.") % {
                    'name': facilitator.name, 'role': role.upper()}))
            entry = by_facilitator.setdefault(facilitator.pk, {'facilitator': facilitator, 'roles': set(), 'villages': [], 'lines': []})
            entry['roles'].add(role.upper())
            if village_id not in entry['villages']:
                entry['villages'].append(village_id)
            entry['lines'].append(number)
        if not people:
            line_errors.append((number, _("No facilitator (FC or FT) on this row.")))
    return by_facilitator, line_errors, line_warnings


def submitted_localities(mode, replace, villages, change_without_submission):
    """Liste à soumettre au calcul de la page « Localités » pour la section importée : les villages du fichier,
    ajoutés à ceux déjà enregistrés sauf en mode remplacement."""
    if mode == MODE_INITIAL:
        current = [] if replace else list(change_without_submission.old_assignment)
    else:
        current = [] if replace else list(change_without_submission.old_stabilization_choices)
    for village_id in villages:
        if village_id not in current:
            current.append(village_id)
    return current
